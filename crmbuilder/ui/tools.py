"""Getting information out, and keeping it sound: exports to Excel and CSV,
the printable record sheet, backups, and the check of everything stored."""
from __future__ import annotations

import csv
import datetime as _dt
import html
import os
import pathlib
import re
import sqlite3
import tkinter as tk
import webbrowser
from tkinter import ttk

from .. import blueprint as bpm
from .. import db as dbm
from .. import sheets, validate
from . import styles, widgets
from .base import Page
from .styles import fs

BIG_JOB = 1500            # records; above this a progress window is shown
MAX_FINDINGS = 500        # rows drawn on the "Check the data" page
DASH = "–"
_BAD_FILE_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]+")
_BAD_CELL_CHARS = re.compile(r"[\000-\010\013\014\016-\037]")
FMT_DATE = "DD/MM/YYYY"
FMT_STAMP = "DD/MM/YYYY HH:MM"
FMT_MONEY = '"£"#,##0.00;-"£"#,##0.00'
FMT_PERCENT = 'General"%"'


def _s(n: int, one: str, many: str) -> str:
    return one if n == 1 else many


def safe_filename(name: str, fallback: str = "Export") -> str:
    name = _BAD_FILE_CHARS.sub(" ", name or "")
    name = " ".join(name.split()).strip(". ")[:120].strip()
    return name or fallback


def _today() -> str:
    return _dt.date.today().isoformat()


def _local(stamp) -> _dt.datetime | None:
    """A stored time (UTC text) as a local date and time."""
    try:
        return _dt.datetime.strptime(str(stamp)[:19], "%Y-%m-%dT%H:%M:%S").replace(
            tzinfo=_dt.timezone.utc).astimezone().replace(tzinfo=None)
    except (ValueError, TypeError):
        return None


def _save_problem(app, exc: Exception):
    widgets.error(app, "The file could not be saved",
                  "If a file with that name is open in Excel or another program, close it there "
                  "and try again, or save under a different name.\n\n"
                  f"({exc})")


class Progress:
    """A small window with a progress bar, shown only for big jobs."""

    def __init__(self, app, title: str, total: int):
        self.app, self.total, self.done, self.win = app, max(1, total), 0, None
        if total < BIG_JOB:
            return
        c = app.c
        win = tk.Toplevel(app.root)
        win.title(title)
        win.configure(bg=c["bg"])
        win.transient(app.root)
        win.resizable(False, False)
        win.protocol("WM_DELETE_WINDOW", lambda: None)
        box = ttk.Frame(win, padding=22)
        box.pack()
        ttk.Label(box, text=title, style="H3.TLabel").pack(anchor="w")
        self.bar = ttk.Progressbar(box, length=fs(340), mode="determinate", maximum=100)
        self.bar.pack(pady=(12, 0))
        self.label = ttk.Label(box, text="", style="Help.TLabel")
        self.label.pack(anchor="w", pady=(8, 0))
        win.update_idletasks()
        root = app.root
        x = root.winfo_rootx() + (root.winfo_width() - win.winfo_reqwidth()) // 2
        y = root.winfo_rooty() + (root.winfo_height() - win.winfo_reqheight()) // 3
        win.geometry(f"+{max(0, x)}+{max(0, y)}")
        self.win = win
        try:
            win.grab_set()
            win.update()
        except tk.TclError:
            pass

    def step(self, n: int = 1, text: str | None = None):
        self.done += n
        if self.win is None:
            return
        if text is not None or self.done % 200 < n:
            try:
                if text is not None:
                    self.label.configure(text=text)
                self.bar.configure(value=100 * min(self.done, self.total) / self.total)
                self.win.update_idletasks()
            except tk.TclError:
                pass

    def close(self):
        if self.win is not None:
            try:
                self.win.destroy()
            except tk.TclError:
                pass
            self.win = None


# ------------------------------------------------------------- exporting
def export_fields(type_def: dict) -> list[dict]:
    """The list's fields in the order the form shows them."""
    return [f for _section, fields in bpm.sections(type_def) for f in fields]


def export_headers(type_def: dict) -> list[str]:
    return ["Ref"] + [f["name"] for f in export_fields(type_def)] + ["Added", "Last changed"]


def cell_text(db, field: dict, value) -> str:
    """A stored value as plain text for a CSV file (numbers without £ or commas,
    dates as dd/mm/yyyy), so that other programs and the importer read it back."""
    if value is None or value == "" or value == []:
        return ""
    kind = field.get("kind")
    if kind in ("link", "user"):
        return validate.display(field, value, db.lookup)
    if kind in ("number", "money", "percent") and isinstance(value, (int, float)) \
            and not isinstance(value, bool):
        if kind == "money":
            return f"{float(value):.2f}"
        return repr(value) if not (isinstance(value, float) and value.is_integer()) else str(int(value))
    return validate.to_edit(field, value)


def cell_value(db, field: dict, value):
    """A stored value as an Excel cell: (value, number format or None)."""
    if value is None or value == "" or value == []:
        return None, None
    kind = field.get("kind")
    try:
        if kind == "date":
            return _dt.date.fromisoformat(str(value)[:10]), FMT_DATE
        if kind == "money":
            return float(value), FMT_MONEY
        if kind == "percent":
            v = float(value)
            return (int(v) if v.is_integer() else v), FMT_PERCENT
        if kind == "number":
            if isinstance(value, bool):
                return str(value), None
            if isinstance(value, (int, float)):
                return value, None
            return float(value), None
    except (TypeError, ValueError):
        return str(value), None
    return cell_text(db, field, value), None


def _put(ws, row: int, col: int, value, fmt=None):
    """Write one cell. Text always stays text (never a formula)."""
    if value is None:
        return
    if isinstance(value, str):
        value = _BAD_CELL_CHARS.sub("", value)
        if not value:
            return
        cell = ws.cell(row=row, column=col, value=value)
        if cell.data_type != "s":
            cell.data_type = "s"
    else:
        cell = ws.cell(row=row, column=col, value=value)
    if fmt:
        cell.number_format = fmt
    return cell


def _finish_sheet(ws, widths: list[int], wrap_cols=()):
    """Bold frozen headings, a filter on every column, sensible widths."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    last = get_column_letter(max(1, len(widths)))
    ws.auto_filter.ref = f"A1:{last}{max(1, ws.max_row)}"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, w + 3))
    if wrap_cols:
        wrap = Alignment(wrap_text=True, vertical="top")
        for col in wrap_cols:
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                if cell.value is not None:
                    cell.alignment = wrap


def _width(value, fmt=None) -> int:
    if value is None:
        return 0
    if fmt == FMT_DATE:
        return 10
    if fmt == FMT_STAMP:
        return 16
    if fmt == FMT_MONEY:
        return len(f"{value:,.2f}") + 1
    text = str(value)
    return max(len(line) for line in text.split("\n")) if text else 0


def fill_record_sheet(ws, db, type_def: dict, records: list[dict], progress: Progress | None = None):
    """Write one list's records to a worksheet."""
    fields = export_fields(type_def)
    headers = export_headers(type_def)
    widths = [len(h) for h in headers]
    for col, h in enumerate(headers, 1):
        _put(ws, 1, col, h)
    last = len(headers)
    for n, r in enumerate(records):
        row = n + 2
        _put(ws, row, 1, r["ref"])
        widths[0] = max(widths[0], len(r["ref"] or ""))
        data = r["data"]
        for i, f in enumerate(fields):
            value, fmt = cell_value(db, f, data.get(f["key"]))
            if value is None:
                continue
            _put(ws, row, i + 2, value, fmt)
            if n < 2000:
                widths[i + 1] = max(widths[i + 1], _width(value, fmt))
        _put(ws, row, last - 1, _local(r.get("created_at")), FMT_STAMP)
        _put(ws, row, last, _local(r.get("updated_at")), FMT_STAMP)
        if progress is not None:
            progress.step()
    widths[last - 2] = widths[last - 1] = 16
    wrap = [i + 2 for i, f in enumerate(fields) if f["kind"] == "longtext"]
    for i in wrap:
        widths[i - 1] = min(widths[i - 1], 50)
    _finish_sheet(ws, widths, wrap)


def write_records_xlsx(db, type_def: dict, records: list[dict], path: str,
                       progress: Progress | None = None):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title(type_def["plural"], [])
    fill_record_sheet(ws, db, type_def, records, progress)
    _save_workbook(wb, path)


def _save_workbook(wb, path: str):
    tmp = path + ".part"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def record_rows(db, type_def: dict, records: list[dict], progress: Progress | None = None):
    """Rows of text for a CSV file, headings first."""
    fields = export_fields(type_def)
    yield export_headers(type_def)
    for r in records:
        data = r["data"]
        yield ([r["ref"]] + [cell_text(db, f, data.get(f["key"])) for f in fields]
               + [validate.format_stamp(r.get("created_at")) if r.get("created_at") else "",
                  validate.format_stamp(r.get("updated_at")) if r.get("updated_at") else ""])
        if progress is not None:
            progress.step()


def write_csv(path: str, rows):
    """UTF-8 with a byte-order mark, which is what makes Excel show £ and accents properly."""
    tmp = path + ".part"
    try:
        with open(tmp, "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(rows)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def sheet_title(name: str, taken: list[str]) -> str:
    """A worksheet name Excel accepts (31 characters, no []:*?/\\), not used yet."""
    base = " ".join(_BAD_SHEET_CHARS.sub(" ", name or "").split()).strip("'")[:31].strip() or "Sheet"
    title, n = base, 2
    low = [t.lower() for t in taken]
    while title.lower() in low:
        tail = f" ({n})"
        title = base[:31 - len(tail)].rstrip() + tail
        n += 1
    return title


def export_records(app, type_def: dict, records: list[dict], fmt: str = "xlsx") -> str | None:
    """Save these records of one list as an Excel workbook or a CSV file."""
    db = app.db
    fmt = "csv" if str(fmt).lower() == "csv" else "xlsx"
    if fmt == "xlsx" and not sheets.have_excel():
        if not widgets.confirm(app, "Excel files cannot be written",
                               "This copy of the app was built without the part that writes Excel "
                               "files. A CSV file holds the same information, and Excel opens it.",
                               yes="Save a CSV file instead"):
            return None
        fmt = "csv"
    name = f"{safe_filename(type_def['plural'])} {_today()}.{fmt}"
    types = [("Excel workbook", "*.xlsx")] if fmt == "xlsx" else [("CSV file", "*.csv")]
    path = widgets.ask_save_file(app, f"Export {type_def['plural'].lower()}", name, types,
                                 defaultextension="." + fmt)
    if not path:
        return None
    progress = Progress(app, f"Exporting {type_def['plural'].lower()}…", len(records))
    try:
        if fmt == "xlsx":
            write_records_xlsx(db, type_def, records, path, progress)
        else:
            write_csv(path, record_rows(db, type_def, records, progress))
    except OSError as exc:
        progress.close()
        _save_problem(app, exc)
        return None
    finally:
        progress.close()
    n = len(records)
    what = (type_def["name"] if n == 1 else type_def["plural"]).lower()
    app.toast(f"Exported {n:,} {what} to {path}", "good")
    return path


NOTE_HEADERS = ["Record ref", "Record", "List", "Kind", "When", "Written by", "Note"]
TASK_HEADERS = ["Record ref", "Record", "List", "Task", "Details", "Due", "For", "Done", "Done on"]


def _everything(db):
    """[(list, its records)], all notes and all tasks, as plain values."""
    lists = [(t, db.records(t["key"])) for t in bpm.active_types(db.blueprint)]
    for _t, recs in lists:
        recs.reverse()                      # oldest first, in reference order
    notes = []
    for t, recs in lists:
        for r in recs:
            for n in reversed(db.notes(r["id"])):
                notes.append((r["ref"], r["title"], t["name"], n["kind"], n["at"],
                              db.user_name(n["created_by"]), n["body"]))
    names = {t["key"]: t["name"] for t, _r in lists}
    tasks = []
    for k in db.tasks(include_done=True):
        if k.get("record_id") and k.get("record_type") not in names:
            continue                         # belongs to a list that was removed from the design
        tasks.append((k.get("record_ref") or "", k.get("record_title") or "",
                      names.get(k.get("record_type"), ""), k["title"], k.get("detail") or "",
                      k.get("due") or "", db.user_name(k.get("assigned_to")),
                      "Yes" if k.get("done_at") else "No", k.get("done_at") or ""))
    return lists, notes, tasks


def _date_or_text(value):
    try:
        return _dt.date.fromisoformat(str(value)[:10]), FMT_DATE
    except ValueError:
        return str(value), None


def write_everything_xlsx(db, path: str, progress: Progress | None = None, data=None) -> dict:
    import openpyxl
    lists, notes, tasks = data or _everything(db)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    titles: list[str] = []
    # the two fixed sheets keep their plain names; a list called "Notes" gets "Notes (2)"
    reserved = ["Notes", "Tasks"]
    for t, recs in lists:
        if progress is not None:
            progress.step(0, f"{t['plural']}…")
        title = sheet_title(t["plural"], titles + reserved)
        titles.append(title)
        fill_record_sheet(wb.create_sheet(title), db, t, recs, progress)

    ws = wb.create_sheet("Notes")
    widths = [len(h) for h in NOTE_HEADERS]
    for col, h in enumerate(NOTE_HEADERS, 1):
        _put(ws, 1, col, h)
    for i, row in enumerate(notes, 2):
        for col, value in enumerate(row, 1):
            fmt = None
            if col == 5:
                when = _local(value)
                value, fmt = (when, FMT_STAMP) if when else (str(value or ""), None)
            _put(ws, i, col, value, fmt)
            if i < 2000:
                widths[col - 1] = max(widths[col - 1], _width(value, fmt))
    widths[6] = min(widths[6], 70)
    _finish_sheet(ws, widths, wrap_cols=[7])

    ws = wb.create_sheet("Tasks")
    widths = [len(h) for h in TASK_HEADERS]
    for col, h in enumerate(TASK_HEADERS, 1):
        _put(ws, 1, col, h)
    for i, row in enumerate(tasks, 2):
        for col, value in enumerate(row, 1):
            fmt = None
            if col == 6 and value:
                value, fmt = _date_or_text(value)
            elif col == 9 and value:
                when = _local(value)
                value, fmt = (when, FMT_STAMP) if when else (str(value), None)
            _put(ws, i, col, value or None, fmt)
            if i < 2000:
                widths[col - 1] = max(widths[col - 1], _width(value, fmt))
    widths[4] = min(widths[4], 50)
    _finish_sheet(ws, widths, wrap_cols=[5])
    if progress is not None:
        progress.step(0, "Saving the file…")
    _save_workbook(wb, path)
    return {"records": sum(len(r) for _t, r in lists), "lists": len(lists),
            "notes": len(notes), "tasks": len(tasks)}


def write_everything_csv(db, folder: str, progress: Progress | None = None, data=None) -> dict:
    """One CSV file per list, plus Notes.csv and Tasks.csv, in folder."""
    lists, notes, tasks = data or _everything(db)
    os.makedirs(folder, exist_ok=True)
    used = ["notes", "tasks"]
    for t, recs in lists:
        name = safe_filename(t["plural"], "List")
        base, n = name, 2
        while name.lower() in used:
            name = f"{base} ({n})"
            n += 1
        used.append(name.lower())
        if progress is not None:
            progress.step(0, f"{t['plural']}…")
        write_csv(os.path.join(folder, name + ".csv"), record_rows(db, t, recs, progress))
    write_csv(os.path.join(folder, "Notes.csv"),
              [NOTE_HEADERS] + [list(r[:4]) + [validate.format_stamp(r[4])] + list(r[5:]) for r in notes])
    write_csv(os.path.join(folder, "Tasks.csv"),
              [TASK_HEADERS] + [list(r[:5]) + [validate.iso_to_uk(r[5]) if r[5] else "", r[6], r[7],
                                               validate.format_stamp(r[8]) if r[8] else ""]
                                for r in tasks])
    return {"records": sum(len(r) for _t, r in lists), "lists": len(lists),
            "notes": len(notes), "tasks": len(tasks)}


def export_all(app) -> str | None:
    """Everything in one Excel workbook: a sheet for each list, plus Notes and
    Tasks. Without Excel support: a folder of CSV files."""
    db = app.db
    stem = f"{safe_filename(db.name, 'CRM')} everything {_today()}"
    excel = sheets.have_excel()
    if excel:
        path = widgets.ask_save_file(app, "Export everything", stem + ".xlsx",
                                     [("Excel workbook", "*.xlsx")], defaultextension=".xlsx")
    else:
        if not widgets.confirm(app, "Excel files cannot be written",
                               "This copy of the app was built without the part that writes Excel "
                               "files. Everything can be saved as CSV files instead – one for "
                               "each list, plus notes and tasks – together in a new folder. "
                               "Excel opens CSV files.",
                               yes="Choose where to put the folder"):
            return None
        folder = widgets.ask_folder(app, "Choose where to put the folder of CSV files")
        path = os.path.join(folder, stem) if folder else ""
        n = 2
        while path and os.path.exists(path):
            path = os.path.join(folder, f"{stem} ({n})")
            n += 1
    if not path:
        return None
    total = sum(db.counts().values())
    progress = Progress(app, "Exporting everything…", total * 2)
    try:
        progress.step(0, "Gathering notes and tasks…")
        data = _everything(db)
        progress.step(total)
        if excel:
            got = write_everything_xlsx(db, path, progress, data)
        else:
            got = write_everything_csv(db, path, progress, data)
    except OSError as exc:
        progress.close()
        _save_problem(app, exc)
        return None
    finally:
        progress.close()
    n = got["records"]
    app.toast(f"Exported {n:,} {_s(n, 'record', 'records')} from {got['lists']} "
              f"{_s(got['lists'], 'list', 'lists')}, with notes and tasks, to {path}", "good")
    return path


# ---------------------------------------------------------- record sheet
SHEET_CSS = """
*{box-sizing:border-box}
html{background:#eef0f4}
body{font-family:"Segoe UI",-apple-system,"Helvetica Neue",Arial,sans-serif;font-size:14px;
line-height:1.45;color:#14171c;background:#eef0f4;margin:0;padding:28px 16px}
.page{max-width:820px;margin:0 auto;background:#fff;border:1px solid #d5dae3;border-radius:6px;
padding:34px 40px}
.top{display:flex;justify-content:space-between;align-items:flex-start;gap:20px;
border-bottom:2px solid #14171c;padding-bottom:14px;margin-bottom:6px}
.kind{font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:#596273;margin:0 0 4px}
h1{font-size:26px;line-height:1.2;margin:0;overflow-wrap:anywhere}
.ref{font-size:13px;color:#596273;text-align:right;white-space:nowrap}
.ref b{display:block;font-size:16px;color:#14171c}
.flag{display:inline-block;margin-top:8px;padding:2px 8px;border:1px solid #b42318;color:#b42318;
border-radius:3px;font-size:12px}
h2{font-size:13px;letter-spacing:.06em;text-transform:uppercase;color:#596273;margin:26px 0 8px;
padding-bottom:4px;border-bottom:1px solid #d5dae3}
dl{display:grid;grid-template-columns:1fr 1fr;gap:10px 28px;margin:0}
dl div{break-inside:avoid;min-width:0}
dl div.wide{grid-column:1 / -1}
dt{font-size:12px;color:#596273}
dd{margin:1px 0 0;white-space:pre-wrap;overflow-wrap:anywhere}
dd.none{color:#98a2b3}
.item{padding:8px 0;border-bottom:1px solid #eceff3;break-inside:avoid}
.item:last-child{border-bottom:0}
.meta{font-size:12px;color:#596273}
.text{white-space:pre-wrap;overflow-wrap:anywhere}
.empty{color:#98a2b3}
ul{margin:0;padding-left:18px}
li{margin:3px 0;overflow-wrap:anywhere}
.foot{margin-top:30px;font-size:12px;color:#596273;border-top:1px solid #d5dae3;padding-top:10px}
.bar{max-width:820px;margin:0 auto 10px;text-align:right;font-size:13px;color:#596273}
.print{font:inherit;font-size:13px;padding:6px 16px;border:1px solid #a9b2c0;border-radius:4px;
background:#fff;color:#14171c;cursor:pointer;margin-left:10px}
@media (max-width:600px){dl{grid-template-columns:1fr}.page{padding:22px 18px}}
@page{margin:16mm}
@media print{html,body{background:#fff;padding:0}.page{border:0;border-radius:0;padding:0;
max-width:none}.bar{display:none}h2{break-after:avoid}}
"""


def record_sheet_html(db, record_id: int) -> str | None:
    """One record as a self-contained web page, laid out for printing."""
    rec = db.get_record(record_id)
    if rec is None:
        return None
    t = bpm.get_type(db.blueprint, rec["type"])
    if t is None:
        return None
    e = lambda s: html.escape("" if s is None else str(s), quote=True)   # noqa: E731
    title = rec["title"] or rec["ref"]
    out = ["<!DOCTYPE html>", '<html lang="en-GB"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           f"<title>{e(title)} – {e(rec['ref'])}</title>", f"<style>{SHEET_CSS}</style></head><body>",
           '<div class="bar">To put this on paper, or save it as a PDF, print this page.'
           '<button class="print" type="button" onclick="window.print()">Print</button></div>',
           '<div class="page">',
           '<div class="top"><div>', f'<p class="kind">{e(t["name"])}</p>', f"<h1>{e(title)}</h1>"]
    if rec.get("deleted_at"):
        out.append('<span class="flag">In the recycle bin</span>')
    out.append(f'</div><div class="ref">Reference<b>{e(rec["ref"])}</b></div></div>')

    for section, fields in bpm.sections(t):
        out.append(f"<h2>{e(section or 'Details')}</h2><dl>")
        for f in fields:
            text = validate.display(f, rec["data"].get(f["key"]), db.lookup)
            wide = ' class="wide"' if f["kind"] in ("longtext", "tags") else ""
            if text:
                out.append(f"<div{wide}><dt>{e(f['name'])}</dt><dd>{e(text)}</dd></div>")
            else:
                out.append(f'<div{wide}><dt>{e(f["name"])}</dt><dd class="none">{DASH}</dd></div>')
        out.append("</dl>")

    notes = db.notes(record_id)
    out.append(f"<h2>Notes ({len(notes)})</h2>")
    if not notes:
        out.append('<p class="empty">No notes.</p>')
    for n in notes:
        meta = [n["kind"], validate.format_stamp(n["at"])]
        who = db.user_name(n["created_by"])
        if who:
            meta.append(who)
        if n.get("edited_at"):
            meta.append("edited")
        out.append(f'<div class="item"><div class="meta">{e("  ·  ".join(meta))}</div>'
                   f'<div class="text">{e(n["body"])}</div></div>')

    tasks = db.tasks(record_id=record_id)
    out.append(f"<h2>Open tasks ({len(tasks)})</h2>")
    if not tasks:
        out.append('<p class="empty">No open tasks.</p>')
    for k in tasks:
        meta = ["Due " + validate.iso_to_uk(k["due"]) if k.get("due") else "No date"]
        who = db.user_name(k.get("assigned_to"))
        if who:
            meta.append("for " + who)
        detail = f'<div class="text">{e(k["detail"])}</div>' if k.get("detail") else ""
        out.append(f'<div class="item"><div><b>{e(k["title"])}</b></div>'
                   f'<div class="meta">{e("  ·  ".join(meta))}</div>{detail}</div>')

    files = db.files(record_id)
    out.append(f"<h2>Files ({len(files)})</h2>")
    if not files:
        out.append('<p class="empty">No files attached.</p>')
    else:
        out.append("<ul>")
        for f in files:
            out.append(f"<li>{e(f['name'])} <span class=\"meta\">({e(_size(f.get('size') or 0))}, "
                       f"added {e(validate.iso_to_uk(f.get('added_at')))})</span></li>")
        out.append("</ul>")

    groups = [(gt, gf, recs) for gt, gf, recs in db.related(record_id) if recs]
    out.append("<h2>Related records</h2>")
    if not groups:
        out.append('<p class="empty">Nothing links to this record.</p>')
    for gt, gf, recs in groups:
        out.append(f'<div class="item"><div class="meta">{e(gt["plural"])} '
                   f'(through {e(gf["name"])})</div><ul>')
        for r in recs:
            out.append(f"<li>{e(r['title'])} <span class=\"meta\">{e(r['ref'])}</span></li>")
        out.append("</ul></div>")

    who = db.user["name"] if db.user else ""
    stamp = _dt.datetime.now().strftime("%d/%m/%Y %H:%M")
    out.append(f'<div class="foot">From {e(db.name)}, {e(stamp)}'
               + (f", by {e(who)}" if who and db.mode == "team" else "") + ".</div>")
    out.append("</div></body></html>")
    return "\n".join(out)


def _size(n: int) -> str:
    for unit in ("bytes", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:,.0f} {unit}" if unit == "bytes" else f"{n:.1f} {unit}"
        n /= 1024
    return ""


def record_sheet(app, record_id: int) -> str | None:
    """Save a printable summary of one record and open it in the web browser."""
    db = app.db
    rec = db.get_record(record_id)
    page = record_sheet_html(db, record_id) if rec else None
    if page is None:
        widgets.error(app, "That did not work", "This record no longer exists.")
        return None
    name = safe_filename(f"{rec['title']} {rec['ref']}", rec["ref"]) + ".html"
    path = widgets.ask_save_file(app, "Save the record sheet", name, [("Web page", "*.html")],
                                 defaultextension=".html")
    if not path:
        return None
    try:
        with open(path, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(page)
    except OSError as exc:
        _save_problem(app, exc)
        return None
    url = pathlib.Path(os.path.abspath(path)).as_uri()
    if app.testing:
        app.test_log.append(("browser", url, ""))
    else:
        try:
            webbrowser.open(url)
        except Exception:
            pass
    app.toast(f"Record sheet saved to {path}. Print it from your web browser.", "good")
    return path


# ---------------------------------------------------------------- backup
def backup_now(app) -> str | None:
    db = app.db
    name = f"{safe_filename(db.name, 'CRM')} backup {_today()}.crm"
    path = widgets.ask_save_file(app, "Back up this CRM", name, [("CRM Builder file", "*.crm")],
                                 defaultextension=".crm")
    if not path:
        return None
    try:
        dest = db.backup_to(path)
    except dbm.DBError as exc:
        widgets.error(app, "The backup was not made", str(exc))
        return None
    except (OSError, sqlite3.Error) as exc:
        widgets.error(app, "The backup was not made",
                      "The copy could not be written there. Check the folder still exists and "
                      f"that you are allowed to save in it, or choose another place.\n\n({exc})")
        return None
    text = f"Backup saved to {dest}"
    if db.encrypted:
        text += (f". Its key file ({os.path.basename(dest)}.keys) was saved beside it – keep the "
                 "two together, the backup cannot be opened without it.")
    app.toast(text, "good")
    return dest


# -------------------------------------------------------- check the data
def _shown(db, field: dict, value) -> str:
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if field.get("kind") in ("link", "user") and isinstance(value, int) and not isinstance(value, bool):
        return db.lookup(field["kind"], value) or f"(number {value})"
    return " ".join(str(value).split())


def check_data(db, progress=None) -> list[dict]:
    """Test every stored value of every list against the current design.

    Returns findings, each {"type", "record", "field", "value" (text), "problem",
    "tidy" (bool), "new" (the tidied value when tidy)}. tidy means the value is
    fine but stored untidily (a phone number without its spaces, say)."""
    out: list[dict] = []
    types = bpm.active_types(db.blueprint)
    user_ids = {u["id"] for u in db.users()}
    targets: dict[int, dict | None] = {}
    total = sum(db.counts().get(t["key"], 0) for t in types)
    done = 0

    def target(rid):
        if rid not in targets:
            targets[rid] = db.get_record(rid)
        return targets[rid]

    for t in types:
        fields = bpm.active_fields(t)
        records = db.records(t["key"])
        records.reverse()
        uniques = {f["key"]: {} for f in fields if f.get("unique")}
        for r in records:
            if progress and done % 200 == 0:
                progress(done, total)
            done += 1
            data = r["data"]

            def found(f, value, problem, tidy=False, new=None, r=r):
                out.append({"type": t, "record": r, "field": f, "value": _shown(db, f, value)
                            if value is not None else "", "problem": problem, "tidy": tidy, "new": new})

            for f in fields:
                value = data.get(f["key"])
                kind = f["kind"]
                if value is None or value == "" or value == []:
                    if f.get("required"):
                        found(f, None, f"{f['name']} is needed, but it is empty.")
                    continue
                if f["key"] in uniques:
                    uniques[f["key"]].setdefault(str(value).strip().lower(), []).append((r, value))
                if kind == "link":
                    if not isinstance(value, int) or isinstance(value, bool):
                        found(f, value, "This should be a link to a record, but it is plain text.")
                        continue
                    other = target(value)
                    if other is None:
                        found(f, value, "The record this pointed to no longer exists.")
                    elif other["type"] != f.get("link_type"):
                        found(f, value, "This points to a record in a different list from the one "
                                        "the field links to now.")
                    elif other["deleted_at"]:
                        found(f, value, "The record this points to is in the recycle bin.")
                    continue
                if kind == "user":
                    if not isinstance(value, int) or isinstance(value, bool) or value not in user_ids:
                        found(f, value, "This should be one of the team, but it does not match anyone.")
                    continue
                if kind == "choice":
                    opts = validate.options_of(f)
                    if str(value) in opts:
                        continue
                    match = next((o for o in opts if o.lower() == str(value).strip().lower()), None)
                    if match is None:
                        found(f, value, f"“{_shown(db, f, value)}” is not one of the choices for "
                                        f"{f['name']} (any more).")
                    else:
                        found(f, value, f"Would be tidied to “{match}”.", True, match)
                    continue
                if kind == "tags":
                    items = value if isinstance(value, (list, tuple)) else \
                        [p.strip() for p in re.split(r"[;,|\n]", str(value)) if p.strip()]
                    new, err = validate.normalise(dict(f, required=False), list(items))
                    if err:
                        opts = [o.lower() for o in validate.options_of(f)]
                        bad = [str(i) for i in items if str(i).strip().lower() not in opts]
                        found(f, value, f"Not among the choices for {f['name']} (any more): "
                              + ", ".join(bad) + ".")
                    elif new != value:
                        found(f, value, f"Would be tidied to “{', '.join(new or [])}”.", True, new)
                    continue
                if kind in ("number", "money", "percent") and isinstance(value, (int, float)) \
                        and not isinstance(value, bool):
                    text = repr(value)
                else:
                    text = validate.to_edit(f, value)
                rule = dict(f, required=False)
                if rule.get("date_rule") == "future":
                    rule.pop("date_rule")       # a date that was ahead when typed may have passed since
                new, err = validate.normalise(rule, text)
                if err:
                    found(f, value, err)
                elif new is not None and (new != value or type(new) is not type(value)
                                          and not (isinstance(new, (int, float))
                                                   and isinstance(value, (int, float))
                                                   and not isinstance(value, bool))):
                    found(f, value, f"Would be tidied to “{validate.display(f, new, db.lookup)}”.",
                          True, new)
        for f in fields:
            for _low, hits in uniques.get(f["key"], {}).items():
                if len(hits) > 1:
                    for r, value in hits:
                        out.append({"type": t, "record": r, "field": f, "value": _shown(db, f, value),
                                    "problem": f"{len(hits)} {t['plural'].lower()} have this "
                                               f"{f['name']}, but it must be different on each.",
                                    "tidy": False, "new": None})
    if progress:
        progress(total, total)
    order = {t["key"]: i for i, t in enumerate(types)}
    out.sort(key=lambda x: (order[x["type"]["key"]], x["record"]["seq"] or 0))
    return out


def tidy_data(db, findings: list[dict]) -> int:
    """Rewrite the values that only need tidying. Returns how many were changed."""
    todo = [x for x in findings if x["tidy"]]
    n = 0
    with db.tx():
        for x in todo:
            db.update_record(x["record"]["id"], {x["field"]["key"]: x["new"]})
            n += 1
    return n


class DataCheckPage(Page):
    """Check the data: does everything stored still fit the design?"""

    def __init__(self, parent, app, **_kw):
        super().__init__(parent, app)
        self.findings: list[dict] | None = None
        self.checked = 0
        self.tree = None
        self._rows: dict[str, dict] = {}
        right = self.header("Check the data", "Does everything stored still fit the design?", back=True)
        self.again_btn = ttk.Button(right, text="Check again", command=self.run_check)
        self.again_btn.pack(side="right")
        why = ttk.Label(
            self, justify="left", style="Help.TLabel",
            text="Every value in every list is tested against the design as it is now. Problems can "
                 "appear when the design was changed after the information was typed in (a field "
                 "became a date, a choice was removed, a field became needed), or when the "
                 "information came from an import.")
        why.pack(fill="x", padx=24, pady=(0, 12))
        self.bind("<Configure>", lambda ev: why.configure(wraplength=max(300, ev.width - 52))
                  if ev.widget is self else None, add="+")
        self.body = ttk.Frame(self)
        self.body.pack(fill="both", expand=True, padx=24, pady=(0, 14))
        self._busy = False
        self._first = self.after(40, self.run_check)

    def destroy(self):
        try:
            self.after_cancel(self._first)
        except (tk.TclError, AttributeError, ValueError):
            pass
        super().destroy()

    def _clear(self):
        for w in self.body.winfo_children():
            w.destroy()
        self.tree = None
        self._rows = {}

    def run_check(self):
        if self._busy or not self.winfo_exists():
            return
        self._busy = True
        self._clear()
        self.again_btn.configure(state="disabled")
        box = ttk.Frame(self.body)
        box.pack(pady=50)
        ttk.Label(box, text="Checking…", style="H2.TLabel").pack()
        bar = ttk.Progressbar(box, length=fs(360), mode="determinate", maximum=100)
        bar.pack(pady=(14, 0))
        self.update_idletasks()

        def progress(done, total):
            try:
                bar.configure(value=100 * done / max(1, total))
                bar.update_idletasks()
            except tk.TclError:
                pass
        try:
            self.findings = check_data(self.db, progress)
            self.checked = sum(self.db.counts().get(t["key"], 0)
                               for t in bpm.active_types(self.db.blueprint))
        finally:
            self._busy = False
        if not self.winfo_exists():
            return
        self.again_btn.configure(state="normal")
        self._draw()

    def refresh(self):
        self.run_check()

    def _draw(self):
        self._clear()
        app, c, found = self.app, self.c, self.findings or []
        lists = len(bpm.active_types(self.db.blueprint))
        if not found:
            box = widgets.card(self.body, app, padding=34)
            box.pack(pady=30)
            ttk.Label(box, text="Everything checks out", style="PanelH2.TLabel",
                      foreground=c["good"]).pack()
            n = self.checked
            text = (f"All {n:,} {_s(n, 'record', 'records')} in {lists} {_s(lists, 'list', 'lists')} "
                    f"{_s(n, 'fits', 'fit')} the design." if n else
                    "There are no records yet, so there is nothing to check.")
            ttk.Label(box, text=text, style="PanelDim.TLabel").pack(pady=(6, 0))
            return
        tidy = sum(1 for x in found if x["tidy"])
        need = len(found) - tidy
        bits = []
        if need:
            bits.append(f"{need:,} {_s(need, 'value needs', 'values need')} a look")
        if tidy:
            bits.append(f"{tidy:,} {_s(tidy, 'is', 'are')} fine but untidy")
        ttk.Label(self.body, text="  ·  ".join(bits), style="H2.TLabel").pack(anchor="w")
        top = ttk.Frame(self.body)
        top.pack(fill="x", pady=(8, 10))
        self.tidy_btn = None
        if tidy and self.db.can_edit:
            self.tidy_btn = ttk.Button(
                top, text=f"Tidy {tidy:,} {_s(tidy, 'value', 'values')} automatically",
                style="Accent.TButton", command=self.tidy)
            self.tidy_btn.pack(side="left", padx=(0, 8))
            widgets.Tooltip(self.tidy_btn, "Rewrites them in their tidy form – a phone number gets "
                                           "its spaces, a postcode its capitals. Nothing else is "
                                           "changed, and each change is kept in the history.", app)
        self.save_btn = ttk.Button(top, text="Save the full list to a file…", command=self.save_list)
        self.save_btn.pack(side="left")

        foot = ttk.Frame(self.body)
        foot.pack(side="bottom", fill="x", pady=(6, 0))
        hint = "Double-click a row to open the record and put it right."
        if len(found) > MAX_FINDINGS:
            hint = (f"Showing the first {MAX_FINDINGS:,} of {len(found):,}. Save the full list to a "
                    "file to see them all.   ") + hint
        ttk.Label(foot, text=hint, style="Help.TLabel").pack(side="left")

        shown = found[:MAX_FINDINGS]
        font = styles.named_font(app.root, "TkDefaultFont")

        def wide(texts, least, most, extra=24):
            longest = max((font.measure(str(x)) for x in texts), default=0) + extra
            return max(fs(least), min(fs(most), longest))
        frame, tree = widgets.make_tree(self.body, ("field", "value", "problem"), show="tree headings",
                                        xscroll=True)
        frame.pack(fill="both", expand=True)
        w0 = wide([f"{x['record']['ref']}  {x['record']['title']}" for x in shown], 160, 340, 64)
        tree.heading("#0", text="Record", anchor="w")
        tree.column("#0", width=w0, minwidth=w0, stretch=False)
        widths = {"field": wide([x["field"]["name"] for x in shown], 90, 240),
                  "value": wide([x["value"] for x in shown] + ["What is stored"], 120, 200),
                  "problem": wide([x["problem"] for x in shown], 220, 760)}
        for col, text in (("field", "Field"), ("value", "What is stored"), ("problem", "What is wrong")):
            tree.heading(col, text=text, anchor="w")
            tree.column(col, width=widths[col], minwidth=widths[col], stretch=col == "problem",
                        anchor="w")
        tree.tag_configure("group", font=(c["ui"], fs(10), "bold"))
        tree.tag_configure("tidy", foreground=c["dim"])
        counts: dict[str, int] = {}
        for x in found:
            counts[x["type"]["key"]] = counts.get(x["type"]["key"], 0) + 1
        groups: dict[str, str] = {}
        for x in shown:
            key = x["type"]["key"]
            if key not in groups:
                n = counts[key]
                groups[key] = tree.insert("", "end", text=f"{x['type']['plural']}  ({n:,})", open=True,
                                          tags=("group",))
            r = x["record"]
            iid = tree.insert(groups[key], "end", text=f"{r['ref']}  {r['title']}",
                              values=(x["field"]["name"], x["value"], x["problem"]),
                              tags=("tidy",) if x["tidy"] else ())
            self._rows[iid] = x
        tree.bind("<Double-1>", self._open_click)
        tree.bind("<Return>", lambda _e: self.open_selected())
        self.tree = tree

    def _open_click(self, event):
        row = self.tree.identify_row(event.y)
        if row in self._rows:
            self.app.open_record(self._rows[row]["record"]["id"])
            return "break"
        return None

    def open_selected(self):
        sel = self.tree.selection() if self.tree is not None else ()
        if sel and sel[0] in self._rows:
            self.app.open_record(self._rows[sel[0]]["record"]["id"])

    def tidy(self):
        if not self.findings or not self.db.can_edit or self._busy:
            return
        try:
            n = tidy_data(self.db, self.findings)
        except dbm.DBError as exc:
            widgets.error(self.app, "The values were not tidied", str(exc))
            return
        self.app.toast(f"Tidied {n:,} {_s(n, 'value', 'values')}", "good")
        self.app.refresh_sidebar()
        self.run_check()

    def save_list(self) -> str | None:
        found = self.findings or []
        name = f"{safe_filename(self.db.name, 'CRM')} data check {_today()}.csv"
        path = widgets.ask_save_file(self.app, "Save the list of problems", name,
                                     [("CSV file", "*.csv")], defaultextension=".csv")
        if not path:
            return None
        rows = [["List", "Ref", "Record", "Field", "What is stored", "What is wrong",
                 "Can be tidied automatically"]]
        for x in found:
            rows.append([x["type"]["plural"], x["record"]["ref"], x["record"]["title"],
                         x["field"]["name"], x["value"], x["problem"], "Yes" if x["tidy"] else "No"])
        try:
            write_csv(path, rows)
        except OSError as exc:
            _save_problem(self.app, exc)
            return None
        self.app.toast(f"Saved {len(found):,} {_s(len(found), 'row', 'rows')} to {path}", "good")
        return path
