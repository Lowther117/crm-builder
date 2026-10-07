"""Reports: simple answers to "how many / how much, split by what".

    Show [Deals] by [Stage] counting [number of deals]

A few headline numbers, a bar chart and the same numbers as a table. Click a
bar to open the list filtered to it. Not a BI tool: one list, one split.

The sums are done by build_report(), which knows nothing about the screen, so
they can be checked on their own.
"""
from __future__ import annotations

import csv
import datetime as _dt
import os
import re
import tkinter as tk
import weakref
from tkinter import font as tkfont
from tkinter import ttk

from .. import blueprint as bpm
from .. import validate
from . import widgets
from .base import Page
from .styles import fs

BLANK = "(blank)"
MAX_BARS = 25
MONTHS_SHOWN = 12
ADDED = "month:__added"          # group key for "Month added"
SPLIT_KINDS = ("choice", "tags", "yesno", "user", "link")
FILTER_KINDS = ("choice", "tags", "user", "yesno")     # what the list page can filter by
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


_fonts: "weakref.WeakKeyDictionary" = weakref.WeakKeyDictionary()


def font_for(app, size: int, bold: bool = False) -> tkfont.Font:
    """A font object for measuring and drawing text, made once per app.

    Tk only keeps a font loaded while some widget is using it, and loading one
    can take tens of milliseconds, so each font is pinned by a label that is
    never shown. Without that, redrawing a canvas (which drops its text items
    first) reloads its fonts every time."""
    cache = _fonts.setdefault(app, {})
    key = (app.c["ui"], size, bold)
    got = cache.get(key)
    if got is None:
        font = tkfont.Font(root=app.root, family=app.c["ui"], size=fs(size),
                           weight="bold" if bold else "normal")
        got = cache[key] = [font, None]
    try:
        alive = got[1] is not None and got[1].winfo_exists()
    except tk.TclError:
        alive = False
    if not alive:
        got[1] = tk.Label(app.root, font=got[0])       # never packed: it only pins the font
    return got[0]


# ------------------------------------------------------------ the choices
def group_choices(t: dict) -> list[tuple[str, str]]:
    """[(key, label)] of the ways a list can be split."""
    out = []
    dates = []
    for f in bpm.active_fields(t):
        if f["kind"] in SPLIT_KINDS:
            out.append((f["key"], f["name"]))
        elif f["kind"] == "date":
            dates.append(f)
    for f in dates:
        out.append(("month:" + f["key"], f"{f['name']} – by month"))
        out.append(("year:" + f["key"], f"{f['name']} – by year"))
    out.append((ADDED, "Month added"))
    return out


def measure_choices(t: dict) -> list[tuple[str, str]]:
    """[(key, label)] of the things that can be counted or added up."""
    out = [("count", f"number of {t['plural'].lower()}")]
    fields = [f for f in bpm.active_fields(t) if f["kind"] in ("money", "number", "percent")]
    for f in fields:
        if f["kind"] != "percent":
            out.append(("sum:" + f["key"], f"total of {f['name']}"))
    for f in fields:
        out.append(("avg:" + f["key"], f"average of {f['name']}"))
    return out


def default_group(t: dict) -> str:
    """The board field if there is one, else the first choice field, else month added."""
    b = bpm.board_field(t)
    if b is not None:
        return b["key"]
    for f in bpm.active_fields(t):
        if f["kind"] == "choice":
            return f["key"]
    return ADDED


def local_date(stamp: str) -> str:
    """'2026-10-07T23:30:00' (UTC, as stored) -> the local date 'YYYY-MM-DD'."""
    try:
        t = _dt.datetime.strptime(str(stamp)[:19], "%Y-%m-%dT%H:%M:%S").replace(
            tzinfo=_dt.timezone.utc).astimezone()
        return t.strftime("%Y-%m-%d")
    except (ValueError, TypeError):
        return str(stamp or "")[:10]


def month_label(ym: str) -> str:
    try:
        return f"{MONTH_NAMES[int(ym[5:7]) - 1]} {ym[:4]}"
    except (ValueError, IndexError):
        return ym


def _num(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def make_formatter(how: str, field: dict | None, values: list):
    """A function that turns one of the report's numbers into text."""
    if how == "count" or field is None:
        return lambda v: "–" if v is None else f"{int(round(v)):,}"
    kind = field["kind"]
    if kind == "money":
        whole = all(v is None or abs(v - round(v)) < 0.005 for v in values)

        def money(v):
            if v is None:
                return "–"
            sign = "-£" if v < -0.004 else "£"
            return sign + (f"{abs(v):,.0f}" if whole else f"{abs(v):,.2f}")
        return money

    def number(v):
        if v is None:
            return "–"
        if how == "sum":
            text = validate.display(dict(field, kind="number"), v)
        else:
            text = f"{v:,.2f}".rstrip("0").rstrip(".")
        return text + ("%" if kind == "percent" else "")
    return number


# ---------------------------------------------------------------- the sums
def build_report(db, t: dict, records: list[dict], group: str, measure: str) -> dict:
    """Split records into groups and count / add up / average each one.

    Returns {"title", "group_label", "measure_label", "rows", "total",
    "total_text", "count", "show_share", "note", "how", "field"} where each row
    is {"label", "value", "text", "count", "share" (0..1 or None), "filter"
    (list page filters that show exactly this group, or None), "muted"}.
    """
    groups = dict(group_choices(t))
    measures = dict(measure_choices(t))
    if group not in groups:
        group = default_group(t)
    if measure not in measures:
        measure = "count"
    how, _, mkey = measure.partition(":")
    mfield = bpm.get_field(t, mkey) if mkey else None
    mode, _, gkey = group.partition(":") if ":" in group else ("field", "", group)
    gfield = bpm.get_field(t, gkey) if gkey != "__added" else None
    kind = gfield["kind"] if (gfield is not None and mode == "field") else mode

    # which bucket(s) each record falls in
    buckets: dict = {}

    def add(key, r):
        b = buckets.get(key)
        if b is None:
            b = buckets[key] = {"count": 0, "sum": 0.0, "n": 0}
        b["count"] += 1
        if mfield is not None:
            v = _num(r["data"].get(mkey))
            if v is not None:
                b["sum"] += v
                b["n"] += 1

    all_b = {"count": 0, "sum": 0.0, "n": 0}
    for r in records:
        all_b["count"] += 1
        if mfield is not None:
            v = _num(r["data"].get(mkey))
            if v is not None:
                all_b["sum"] += v
                all_b["n"] += 1
        if mode in ("month", "year"):
            raw = local_date(r.get("created_at")) if gkey == "__added" else r["data"].get(gkey)
            raw = str(raw or "")
            add((raw[:7] if mode == "month" else raw[:4]) or None, r)
        elif kind == "tags":
            tags = r["data"].get(gkey) or []
            if not isinstance(tags, (list, tuple)):
                tags = [tags]
            if not tags:
                add(None, r)
            for tag in dict.fromkeys(tags):
                add(tag, r)
        elif kind == "yesno":
            add(bool(r["data"].get(gkey)), r)
        else:
            v = r["data"].get(gkey)
            add(None if v in (None, "") else v, r)

    def value_of(b):
        if how == "count":
            return b["count"]
        if how == "sum":
            return b["sum"]
        return (b["sum"] / b["n"]) if b["n"] else None

    def size(key):
        v = value_of(buckets[key])
        return -1e300 if v is None else v

    can_filter = kind in FILTER_KINDS
    rows = []

    def row(key, label, keys=None, muted=False):
        if keys is None:
            b = buckets.get(key) or {"count": 0, "sum": 0.0, "n": 0}
            flt = {gkey: [key]} if can_filter else None
        else:                                   # several groups rolled into one
            b = {"count": 0, "sum": 0.0, "n": 0}
            for k in keys:
                for part in ("count", "sum", "n"):
                    b[part] += buckets[k][part]
            flt = {gkey: list(keys)} if can_filter else None
        if b["count"] == 0:
            flt = None
        rows.append({"label": label, "value": value_of(b), "count": b["count"], "filter": flt,
                     "muted": muted, "key": key})

    def other_label(sofar):
        return "All others" if any(x["label"] == "Other" for x in sofar) else "Other"

    present = [k for k in buckets if k is not None]
    has_blank = None in buckets
    room = MAX_BARS - (1 if has_blank else 0)
    if mode in ("month", "year"):
        keys = sorted(present)
        limit = MONTHS_SHOWN if mode == "month" else room - 1
        name = month_label if mode == "month" else (lambda y: y)
        if len(keys) > limit:
            old, keys = keys[:-limit], keys[-limit:]
            row("__before", f"Before {name(keys[0])}", keys=old, muted=True)
            rows[-1]["bar"] = False     # everything older in one lump would dwarf the months
        for k in keys:
            row(k, name(k))
    elif kind == "yesno":
        row(True, "Yes")
        row(False, "No")
        has_blank = False
    elif kind == "choice":
        options = validate.options_of(gfield)
        extra = sorted((k for k in present if k not in options), key=lambda k: (-size(k), str(k)))
        keys = options + extra
        if len(keys) > room:
            keep = set(sorted(keys, key=lambda k: -size(k) if k in buckets else 1e300)[:room - 1])
            other = [k for k in keys if k not in keep and k in buckets]
            keys = [k for k in keys if k in keep]
        else:
            other = []
        for k in keys:
            row(k, str(k))
        if other:
            row("__other", other_label(rows), keys=other, muted=True)
    else:                                       # tags, team member, link: largest first
        def label_of(k):
            if kind in ("user", "link"):
                return db.lookup(kind, k) or f"(record {k})"
            return str(k)
        labelled = sorted(((k, label_of(k)) for k in present),
                          key=lambda kl: (-size(kl[0]), kl[1].lower()))
        other = []
        if len(labelled) > room:
            labelled, rest = labelled[:room - 1], labelled[room - 1:]
            other = [k for k, _l in rest]
        for k, label in labelled:
            row(k, label)
        if other:
            row("__other", other_label(rows), keys=other, muted=True)
    if has_blank:
        row(None, BLANK, muted=True)

    total = value_of(all_b)
    fmt = make_formatter(how, mfield, [r["value"] for r in rows] + [total])
    show_share = how != "avg"
    positive = all((r["value"] or 0) >= 0 for r in rows)
    for r in rows:
        r["text"] = fmt(r["value"])
        r["share"] = None
        if show_share and total and positive and r["value"] is not None:
            r["share"] = r["value"] / float(total)
    note = ""
    if kind == "tags" and any(len(r["data"].get(gkey) or []) > 1 for r in records
                              if isinstance(r["data"].get(gkey), (list, tuple))):
        note = (f"A {t['name'].lower()} with more than one of these is counted under each, "
                "so the rows can add up to more than the total.")
    glabel = groups[group]
    short = glabel.replace(" – by month", " (month)").replace(" – by year", " (year)")
    by = "month added" if group == ADDED else short
    return {"title": f"{t['plural']} by {by}",
            "group": group, "measure": measure, "group_label": short,
            "group_head": {"month": "Month", "year": "Year"}.get(mode, short) if gkey != "__added" else short,
            "measure_label": measures[measure][0].upper() + measures[measure][1:],
            "rows": rows, "total": total, "total_text": fmt(total), "count": all_b["count"],
            "show_share": show_share, "note": note, "how": how, "field": mfield, "kind": kind}


def headline(db, t: dict, records: list[dict]) -> list[tuple[str, str, str]]:
    """Two to four big numbers about a list: [(label, value, small print)]."""
    n = len(records)
    out = [(t["plural"], f"{n:,}", "in this list")]
    today = _dt.date.today()
    this = today.strftime("%Y-%m")
    last = (today.replace(day=1) - _dt.timedelta(days=1)).strftime("%Y-%m")
    added = [local_date(r.get("created_at"))[:7] for r in records]
    out.append(("Added this month", f"{added.count(this):,}", f"{added.count(last):,} last month"))
    money = bpm.money_field(t)
    if money is not None:
        vals = [v for v in (_num(r["data"].get(money["key"])) for r in records) if v is not None]
        total = sum(vals)
        fmt = make_formatter("sum", money, [total])
        small = ""
        if vals:
            mean = float(round(sum(vals) / len(vals)))
            small = "average " + make_formatter("avg", money, [mean])(mean)
        out.append((f"Total {money['name'][0].lower() + money['name'][1:]}", fmt(total), small))
    board = bpm.board_field(t)
    if board is not None and n:
        counts: dict = {}
        for r in records:
            v = r["data"].get(board["key"])
            if v not in (None, ""):
                counts[v] = counts.get(v, 0) + 1
        if counts:
            options = validate.options_of(board)
            top = max(counts, key=lambda k: (counts[k], -options.index(k) if k in options else -999))
            out.append((f"Most common {board['name'][0].lower() + board['name'][1:]}", str(top),
                        f"{counts[top]:,} of {n:,}"))
    return out


def safe_file_name(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', " ", text).strip() or "Report"


def write_report(path: str, report: dict) -> None:
    """Save a report's numbers as .xlsx (when openpyxl is there) or .csv."""
    head = [report["group_label"], report["measure_label"]]
    if report["show_share"]:
        head.append("% of total")
    field = report["field"]
    if path.lower().endswith(".xlsx"):
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.append(head)
        for r in report["rows"]:
            line = [r["label"], r["value"]]
            if report["show_share"]:
                line.append(r["share"])
            ws.append(line)
        last = ["Total" if report["how"] != "avg" else "Overall", report["total"]]
        if report["show_share"]:
            last.append(1 if report["total"] else None)
        ws.append(last)
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold
        for cell in ws[ws.max_row]:
            cell.font = bold
        number = "#,##0"
        if report["how"] != "count" and field is not None:
            number = {"money": "£#,##0.00", "percent": '0.0"%"'}.get(field["kind"], "#,##0.##")
        for row in ws.iter_rows(min_row=2):
            row[1].number_format = number
            if len(row) > 2:
                row[2].number_format = "0.0%"
        ws.column_dimensions["A"].width = max(14, min(60, max(len(str(c.value or "")) for c in ws["A"]) + 2))
        ws.column_dimensions["B"].width = max(14, len(head[1]) + 2)
        ws.column_dimensions["C"].width = 12
        ws.freeze_panes = "A2"
        wb.save(path)
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(head)

        def plain(v):
            if v is None:
                return ""
            return int(v) if float(v).is_integer() else round(v, 2)
        for r in report["rows"]:
            line = [r["label"], plain(r["value"])]
            if report["show_share"]:
                line.append("" if r["share"] is None else f"{r['share'] * 100:.1f}%")
            w.writerow(line)
        last = ["Total" if report["how"] != "avg" else "Overall", plain(report["total"])]
        if report["show_share"]:
            last.append("100%" if report["total"] else "")
        w.writerow(last)


def have_excel() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


# ---------------------------------------------------------------- the page
class ReportsPage(Page):
    nav_key = "reports"
    auto_refresh = True

    def __init__(self, parent, app, type_key=None, group=None, measure=None, **_kw):
        super().__init__(parent, app)
        c = self.c
        self.types = bpm.active_types(self.db.blueprint)
        self.t = bpm.get_type(self.db.blueprint, type_key or "")
        if self.t is None or self.t.get("archived"):
            self.t = self.types[0] if self.types else None
        self.group = group
        self.measure = measure
        self.report = None
        self.records: list[dict] = []
        self._hot = None
        self._redraw = None
        self._wide = None
        self.f_body = font_for(app, 10)
        self.f_bold = font_for(app, 10, True)
        self.f_small = font_for(app, 9)
        self.f_head = font_for(app, 9, True)
        self.row_h = self.f_body.metrics("linespace") + 14
        right = self.header("Reports", "Count up what is in a list, split whichever way is useful.")
        if self.t is None:
            self._no_lists()
            return
        self.export_btn = ttk.Button(right, text="Export these numbers…", command=self.export)
        self.export_btn.pack(side="right")
        self._build_pickers()
        self.scroll = widgets.ScrollFrame(self, app)
        self.scroll.pack(fill="both", expand=True, padx=(24, 0), pady=(0, 6))
        self.area = self.scroll.body
        self._fill_pickers()
        self.load()

    # ------------------------------------------------------------ pickers
    def _build_pickers(self):
        bar = ttk.Frame(self)
        bar.pack(fill="x", padx=24, pady=(2, 14))
        self.type_var, self.group_var, self.measure_var = (tk.StringVar(master=self) for _ in range(3))

        def word(text, col, first=False):
            ttk.Label(bar, text=text).grid(row=0, column=col, padx=(0 if first else 10, 8))

        def box(var, col, handler):
            cb = ttk.Combobox(bar, textvariable=var, state="readonly", exportselection=False)
            cb.grid(row=0, column=col, sticky="ew")
            cb.bind("<<ComboboxSelected>>", handler)
            bar.columnconfigure(col, weight=1)
            return cb
        word("Show", 0, True)
        self.type_box = box(self.type_var, 1, self._type_picked)
        word("by", 2)
        self.group_box = box(self.group_var, 3, self._group_picked)
        word("counting", 4)
        self.measure_box = box(self.measure_var, 5, self._measure_picked)
        bar.columnconfigure(6, weight=1000)      # spare room goes here, not into the boxes
        names = [t["plural"] for t in self.types]
        self.type_box.configure(values=names, width=max(10, min(22, max(len(n) for n in names) + 2)))

    def _fill_pickers(self):
        """Set the three boxes for the current list, keeping valid choices."""
        t = self.t
        self.groups = group_choices(t)
        self.measures = measure_choices(t)
        if self.group not in dict(self.groups):
            self.group = default_group(t)
        if self.measure not in dict(self.measures):
            self.measure = "count"
        self.type_var.set(t["plural"])
        glabels = [label for _k, label in self.groups]
        mlabels = [label for _k, label in self.measures]
        self.group_box.configure(values=glabels, width=max(12, min(30, max(len(x) for x in glabels) + 1)))
        self.measure_box.configure(values=mlabels, width=max(12, min(26, max(len(x) for x in mlabels) + 1)))
        self.group_var.set(dict(self.groups)[self.group])
        self.measure_var.set(dict(self.measures)[self.measure])

    def _type_picked(self, _e=None):
        i = self.type_box.current()
        if i < 0 or self.types[i] is self.t:
            return
        self.t = self.types[i]
        self.group = self.measure = None
        self._fill_pickers()
        self.type_box.selection_clear()
        self.load()

    def _group_picked(self, _e=None):
        i = self.group_box.current()
        if i >= 0:
            self.group = self.groups[i][0]
            self.group_box.selection_clear()
            self.load()

    def _measure_picked(self, _e=None):
        i = self.measure_box.current()
        if i >= 0:
            self.measure = self.measures[i][0]
            self.measure_box.selection_clear()
            self.load()

    def set_choices(self, type_key=None, group=None, measure=None):
        """Change what is shown (what the three boxes do; handy for tests)."""
        if type_key is not None and type_key != self.t["key"]:
            t = bpm.get_type(self.db.blueprint, type_key)
            if t is not None and not t.get("archived"):
                self.t = t
                self.group = self.measure = None
        if group is not None:
            self.group = group
        if measure is not None:
            self.measure = measure
        self._fill_pickers()
        self.load()

    # --------------------------------------------------------------- data
    def load(self):
        """Read the records, do the sums and redraw everything below the boxes."""
        t = self.t
        self.records = self.db.records(t["key"])
        self.report = build_report(self.db, t, self.records, self.group, self.measure)
        self._hot = None
        for w in self.area.winfo_children():
            w.destroy()
        self.chart = self.table = None
        self._wide = None
        if not self.records:
            self.export_btn.state(["disabled"])
            self._no_records()
            return
        self.export_btn.state(["!disabled"])
        self._build_tiles()
        only_added = len(self.groups) == 1
        notes = []
        if only_added:
            notes.append(f"{t['plural']} have no choice, tick-box, date or link fields to split by yet, "
                         "so they are shown by the month they were added.")
        if self.report["note"]:
            notes.append(self.report["note"])
        if notes:
            row = ttk.Frame(self.area)
            row.pack(fill="x", padx=(0, 24), pady=(12, 0))
            lab = ttk.Label(row, text=" ".join(notes), style="Help.TLabel", justify="left")
            lab.pack(side="left", anchor="w")
            row.bind("<Configure>", lambda e, lab=lab: lab.configure(wraplength=max(200, e.width - 8)))
            if only_added and self.db.is_admin:
                link_row = ttk.Frame(self.area)
                link_row.pack(fill="x", padx=(0, 24), pady=(2, 0))
                widgets.link(link_row, "Add a field to split by",
                             lambda: self.app.go("design", type_key=t["key"])).pack(side="left")
        self._build_main()

    def refresh(self):
        if self.t is not None:
            self.load()

    def state(self) -> dict:
        if self.t is None:
            return {}
        return {"type_key": self.t["key"], "group": self.group, "measure": self.measure}

    # ------------------------------------------------------- empty states
    def _no_lists(self):
        box = widgets.card(self, self.app, padding=34)
        box.pack(pady=40)
        ttk.Label(box, text="Nothing to report on yet", style="PanelH2.TLabel").pack()
        ttk.Label(box, text="Reports count up what is in your lists. Make a list first, then come back.",
                  style="PanelDim.TLabel").pack(pady=(6, 16 if self.db.is_admin else 0))
        if self.db.is_admin:
            ttk.Button(box, text="Design your first list", style="Accent.TButton",
                       command=lambda: self.app.go("design")).pack()

    def _no_records(self):
        t = self.t
        box = widgets.card(self.area, self.app, padding=34)
        box.pack(pady=40)
        ttk.Label(box, text=f"No {t['plural'].lower()} yet", style="PanelH2.TLabel").pack()
        ttk.Label(box, text=f"Once there are some {t['plural'].lower()}, this page counts them up for you.",
                  style="PanelDim.TLabel").pack(pady=(6, 16 if self.db.can_edit else 0))
        if self.db.can_edit:
            ttk.Button(box, text=f"Add your first {t['name'].lower()}", style="Accent.TButton",
                       command=lambda: self.app.new_record(t["key"])).pack()

    # -------------------------------------------------------------- tiles
    def _build_tiles(self):
        c = self.c
        self.tiles = ttk.Frame(self.area)
        self.tiles.pack(fill="x", padx=(0, 24))
        self._tile_values = []
        items = headline(self.db, self.t, self.records)
        slot_h = font_for(self.app, 20, True).metrics("linespace") + 2
        for i, (label, value, small) in enumerate(items):
            card = widgets.card(self.tiles, self.app, padding=14)
            card.grid(row=0, column=i, sticky="nsew", padx=(0, 12 if i < len(items) - 1 else 0))
            self.tiles.columnconfigure(i, weight=1, uniform="tile")
            ttk.Label(card, text=label, style="PanelHelp.TLabel").pack(anchor="w")
            # a fixed-height slot, so tiles stay level when a long number has to shrink
            slot = tk.Frame(card, bg=c["panel"], height=slot_h)
            slot.pack(fill="x", pady=(2, 0))
            slot.pack_propagate(False)
            big = tk.Label(slot, text=value, bg=c["panel"], fg=c["text"], anchor="w",
                           font=(c["ui"], fs(20), "bold"), padx=0, pady=0, bd=0)
            big.pack(side="bottom", anchor="w", fill="x")
            ttk.Label(card, text=small or " ", style="PanelHelp.TLabel").pack(anchor="w")
            self._tile_values.append((card, big, value))
        self.tiles.bind("<Configure>", self._fit_tiles)

    def _fit_tiles(self, event=None):
        """Big numbers shrink a step or two rather than being cut off."""
        c = self.c
        n = max(1, len(self._tile_values))
        width = (self.tiles.winfo_width() - 12 * (n - 1)) // n - 2 * 14 - 4
        if width < 40:
            return
        for _card, big, value in self._tile_values:
            text, size = value, 11
            for size in (20, 17, 14, 12, 11):
                if font_for(self.app, size, True).measure(value) <= width:
                    break
            else:
                text = self._fit(font_for(self.app, 11, True), value, width)
            if big.cget("text") != text or getattr(big, "shown_size", None) != size:
                big.shown_size = size
                big.configure(text=text, font=font_for(self.app, size, True))

    # ------------------------------------------------------ chart + table
    def _build_main(self):
        c, rep = self.c, self.report
        self.main = ttk.Frame(self.area)
        self.main.pack(fill="x", padx=(0, 24), pady=(12, 18))
        self.chart_card = widgets.card(self.main, self.app, padding=16)
        ttk.Label(self.chart_card, text=rep["title"], style="PanelH2.TLabel").pack(anchor="w")
        clickable = any(r["filter"] for r in rep["rows"])
        sub = rep["measure_label"] + (".  Click a bar to see what is behind it." if clickable else ".")
        ttk.Label(self.chart_card, text=sub, style="PanelHelp.TLabel").pack(anchor="w", pady=(2, 10))
        self.chart = tk.Canvas(self.chart_card, bg=c["panel"], highlightthickness=1, borderwidth=0,
                               highlightbackground=c["panel"], highlightcolor=c["accent"],
                               height=len(rep["rows"]) * self.row_h + 2, width=200, takefocus=1)
        self.chart.pack(fill="x")
        self.table_card = widgets.card(self.main, self.app, padding=16)
        self.table = tk.Canvas(self.table_card, bg=c["panel"], highlightthickness=0, borderwidth=0,
                               height=(len(rep["rows"]) + 2) * self.row_h + 10, width=fs(10) * 30)
        self.table.pack(fill="x")
        for cv in (self.chart, self.table):
            cv.bind("<Configure>", self._queue_redraw)
            cv.bind("<Motion>", lambda e, cv=cv: self._set_hot(self._row_at(cv, e.y)))
            cv.bind("<Leave>", lambda _e: self._set_hot(None))
            cv.bind("<Button-1>", lambda e, cv=cv: self._click(cv, e))
        ch = self.chart
        ch.bind("<Up>", lambda _e: self._key_step(-1))
        ch.bind("<Down>", lambda _e: self._key_step(1))
        ch.bind("<Return>", lambda _e: self.open_row(self._hot) or "break")
        ch.bind("<space>", lambda _e: self.open_row(self._hot) or "break")
        ch.bind("<FocusIn>", lambda _e: self._set_hot(0 if self._hot is None else self._hot))
        ch.bind("<FocusOut>", lambda _e: self._set_hot(None))
        self.main.bind("<Configure>", self._arrange)
        self._arrange()

    def _arrange(self, event=None):
        """Chart and table side by side when there is room, else stacked."""
        width = event.width if event is not None else self.main.winfo_width()
        if width < 50:
            width = self.app.content.winfo_width() - 60
        wide = width >= fs(10) * 80
        if wide == self._wide:
            return
        self._wide = wide
        main = self.main
        self.chart_card.grid_forget()
        self.table_card.grid_forget()
        if wide:
            main.columnconfigure(0, weight=1)
            main.columnconfigure(1, weight=0)
            self.chart_card.grid(row=0, column=0, sticky="new")
            self.table_card.grid(row=0, column=1, sticky="new", padx=(12, 0))
        else:
            main.columnconfigure(0, weight=1)
            main.columnconfigure(1, weight=0, minsize=0)
            self.chart_card.grid(row=0, column=0, sticky="new")
            self.table_card.grid(row=1, column=0, sticky="new", pady=(12, 0))

    def _queue_redraw(self, _e=None):
        if self._redraw is None:
            self._redraw = self.after_idle(self.draw)

    def draw(self):
        """Draw the bars and the table at the current width."""
        self._redraw = None
        if self.chart is None or not self.chart.winfo_exists():
            return
        self._draw_chart()
        self._draw_table()
        self._paint_hot()

    def _fit(self, font, text, width):
        if font.measure(text) <= width:
            return text
        while len(text) > 1 and font.measure(text + "…") > width:
            text = text[:-1]
        return text.rstrip() + "…"

    def _draw_chart(self):
        cv, c, rep, rh = self.chart, self.c, self.report, self.row_h
        cv.delete("all")
        rows = rep["rows"]
        width = cv.winfo_width()
        if width < 120:
            return
        label_w = min(max([self.f_body.measure(r["label"]) for r in rows] + [20]),
                      int(width * 0.34), fs(10) * 24)
        value_w = max([self.f_body.measure(r["text"]) for r in rows] + [10])
        x_bar = 8 + label_w + 14
        bar_max = max(20, width - x_bar - 10 - value_w - 10)
        biggest = max([abs(r["value"] or 0) for r in rows if r.get("bar", True)] + [0]) or 1
        colour = self.t.get("color") or c["accent"]
        for i, r in enumerate(rows):
            y = 1 + i * rh
            cv.create_rectangle(1, y, width - 1, y + rh, fill=c["panel"], outline="", tags=(f"hot{i}",))
            cv.create_text(8, y + rh // 2, anchor="w", text=self._fit(self.f_body, r["label"], label_w),
                           font=self.f_body, fill=c["dim"] if r["muted"] else c["text"])
            v = r["value"] or 0
            x_end = x_bar
            if not r.get("bar", True):
                x_end = x_bar - 8          # a lump of older dates: the number only, no bar
            elif v:
                x_end = x_bar + max(3, int(round(bar_max * abs(v) / biggest)))
                fill = c["bad"] if v < 0 else (c["field_border"] if r["muted"] else colour)
                cv.create_rectangle(x_bar, y + 7, x_end, y + rh - 7, fill=fill, outline="")
            else:
                cv.create_line(x_bar, y + 7, x_bar, y + rh - 7, fill=c["border"])
            cv.create_text(x_end + 8, y + rh // 2, anchor="w", text=r["text"], font=self.f_body,
                           fill=c["text"] if (v and r.get("bar", True)) else c["dim"])

    def _draw_table(self):
        cv, c, rep, rh = self.table, self.c, self.report, self.row_h
        cv.delete("all")
        rows = rep["rows"]
        width = cv.winfo_width()
        if width < 120:
            return
        share = rep["show_share"]
        x_share = width - 8
        share_w = max(self.f_head.measure("% of total"), self.f_body.measure("100.0%")) if share else 0
        x_value = x_share - (share_w + 18 if share else 0)
        head = rep["measure_label"]
        if rep["how"] != "count" and rep["field"] is not None:
            head = ("Total " if rep["how"] == "sum" else "Average ") + rep["field"]["name"].lower()
        elif rep["how"] == "count":
            head = "Number"
        value_w = max([self.f_body.measure(r["text"]) for r in rows]
                      + [self.f_bold.measure(rep["total_text"]), self.f_head.measure(head)])
        label_w = max(40, x_value - value_w - 16 - 8)
        mid = rh // 2
        cv.create_text(8, mid, anchor="w", text=self._fit(self.f_head, rep["group_head"], label_w),
                       font=self.f_head, fill=c["dim"])
        cv.create_text(x_value, mid, anchor="e", text=head, font=self.f_head, fill=c["dim"])
        if share:
            cv.create_text(x_share, mid, anchor="e", text="% of total", font=self.f_head, fill=c["dim"])
        cv.create_line(0, rh, width, rh, fill=c["border"])
        y = rh + 3
        for i, r in enumerate(rows):
            cv.create_rectangle(0, y, width, y + rh, fill=c["panel"], outline="", tags=(f"hot{i}",))
            tone = c["dim"] if (r["muted"] or not r["count"]) else c["text"]
            cv.create_text(8, y + mid, anchor="w", text=self._fit(self.f_body, r["label"], label_w),
                           font=self.f_body, fill=tone)
            cv.create_text(x_value, y + mid, anchor="e", text=r["text"], font=self.f_body, fill=tone)
            if share:
                cv.create_text(x_share, y + mid, anchor="e", font=self.f_body, fill=c["dim"],
                               text="–" if r["share"] is None else f"{r['share'] * 100:.1f}%")
            y += rh
        y += 3
        cv.create_line(0, y, width, y, fill=c["border"])
        name = "Total" if rep["how"] != "avg" else f"All {self.t['plural'].lower()}"
        cv.create_text(8, y + mid + 1, anchor="w", text=name, font=self.f_bold, fill=c["text"])
        cv.create_text(x_value, y + mid + 1, anchor="e", text=rep["total_text"], font=self.f_bold,
                       fill=c["text"])
        if share and rep["total"] and rep["kind"] != "tags":
            cv.create_text(x_share, y + mid + 1, anchor="e", text="100%", font=self.f_bold, fill=c["dim"])

    # ------------------------------------------------- hover, keys, click
    def _row_at(self, cv, y):
        top = 1 if cv is self.chart else self.row_h + 3
        i = int((y - top) // self.row_h) if y >= top else -1
        return i if 0 <= i < len(self.report["rows"]) else None

    def _set_hot(self, i):
        if i == self._hot:
            return
        self._hot = i
        self._paint_hot()

    def _paint_hot(self):
        c = self.c
        rows = self.report["rows"] if self.report else []
        hot = self._hot
        for cv in (self.chart, self.table):
            if cv is None or not cv.winfo_exists():
                continue
            for i in range(len(rows)):
                cv.itemconfigure(f"hot{i}", fill=c["sel"] if i == hot else c["panel"])
            hand = hot is not None and hot < len(rows) and rows[hot]["filter"]
            cv.configure(cursor="hand2" if hand else "")

    def _key_step(self, step):
        n = len(self.report["rows"])
        if n:
            i = 0 if self._hot is None else max(0, min(n - 1, self._hot + step))
            self._set_hot(i)
            try:
                y = self.chart.winfo_rooty() - self.area.winfo_rooty() + i * self.row_h
                top = self.scroll.canvas.canvasy(0)
                view = self.scroll.canvas.winfo_height()
                total = max(1, self.area.winfo_reqheight())
                if y < top + 10:
                    self.scroll.canvas.yview_moveto(max(0, y - 20) / total)
                elif y + self.row_h > top + view - 10:
                    self.scroll.canvas.yview_moveto((y + self.row_h + 20 - view) / total)
            except tk.TclError:
                pass
        return "break"

    def _click(self, cv, event):
        if cv is self.chart:
            cv.focus_set()
        i = self._row_at(cv, event.y)
        self._set_hot(i)
        self.open_row(i)

    def open_row(self, i) -> bool:
        """Open the list filtered to one bar, where the list page can do that."""
        rows = self.report["rows"] if self.report else []
        if i is None or not (0 <= i < len(rows)) or not rows[i]["filter"]:
            return False
        self.app.go("list", type_key=self.t["key"],
                    filters={k: list(v) for k, v in rows[i]["filter"].items()})
        return True

    # ------------------------------------------------------------- export
    def export(self):
        rep = self.report
        if rep is None or not self.records:
            self.app.toast("There is nothing to export yet.")
            return
        excel = have_excel()
        ext = ".xlsx" if excel else ".csv"
        types = [("Excel workbook", "*.xlsx")] if excel else []
        types.append(("CSV (opens in any spreadsheet)", "*.csv"))
        path = widgets.ask_save_file(self.app, "Export these numbers",
                                     safe_file_name(rep["title"]) + ext, types, defaultextension=ext)
        if not path:
            return
        low = path.lower()
        if not low.endswith((".xlsx", ".csv")):
            path += ext
        elif low.endswith(".xlsx") and not excel:
            path = path[:-5] + ".csv"
        try:
            write_report(path, rep)
        except OSError as exc:
            widgets.error(self.app, "The file could not be saved",
                          f"{os.path.basename(path)} could not be written. If it is open in another "
                          f"program, close it there and try again.\n\n({exc})")
            return
        self.app.toast(f"Saved {os.path.basename(path)}", "good")
