"""Reading spreadsheets (CSV and Excel) into plain rows of text."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import os
import re
from decimal import Decimal


class SheetError(Exception):
    """A problem reading the file, worded for the person who chose it."""


EXCEL = (".xlsx", ".xlsm")
TEXT = (".csv", ".tsv", ".txt", ".tab")


def have_excel() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def file_types() -> list[tuple[str, str]]:
    if have_excel():
        return [("Spreadsheets", "*.xlsx *.xlsm *.csv *.tsv *.txt"), ("Excel workbook", "*.xlsx *.xlsm"),
                ("CSV file", "*.csv *.tsv *.txt"), ("All files", "*.*")]
    return [("CSV file", "*.csv *.tsv *.txt"), ("All files", "*.*")]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, _dt.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime("%d/%m/%Y %H:%M")
        return v.strftime("%d/%m/%Y")
    if isinstance(v, _dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        if v - v != 0:                      # infinity, not-a-number
            return ""
        if v.is_integer():
            return str(int(v))
        # (rounding hides the 0.30000000000000004 kind of noise)
        r = repr(round(v, 10) if abs(v) >= 1e-6 else v)
        if "e" in r:                        # written out in full, never as 1e-07
            r = format(Decimal(r), "f")
        return r.rstrip("0").rstrip(".") if "." in r else r
    return str(v).strip()


_FORMATS: dict[str, str] = {}


def _excel_cell(c) -> str:
    """A cell as the person sees it in Excel, where that differs from what is
    stored: 60% is stored as 0.6, and 00123 as 123 with a format of 00000."""
    v = c.value
    if v is None:
        return ""
    if getattr(c, "data_type", "") == "e":       # #N/A, #DIV/0! ... are not information
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        fmt = str(getattr(c, "number_format", "") or "")
        if fmt not in _FORMATS:     # (a quoted "%" is only a label, not a percentage)
            _FORMATS[fmt] = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", fmt).split(";")[0]
        fmt = _FORMATS[fmt]
        if "%" in fmt:
            return _cell(round(v * 100, 8)) + "%"
        if re.fullmatch(r"0{2,}", fmt) and float(v).is_integer() and v >= 0:
            return str(int(v)).zfill(len(fmt))
    return _cell(v)


def _visible_first(wb) -> list[str]:
    """Worksheet names, the ones Excel shows before any hidden ones."""
    def hidden(name):
        return getattr(wb[name], "sheet_state", "visible") != "visible"
    names = [ws.title for ws in wb.worksheets]
    return [n for n in names if not hidden(n)] + [n for n in names if hidden(n)]


def sheet_names(path: str) -> list[str]:
    """Worksheet names of an Excel file ([] for CSV). Hidden sheets come last."""
    if not path.lower().endswith(EXCEL):
        return []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return _visible_first(wb)
        finally:
            wb.close()
    except ImportError:
        raise SheetError("Excel files cannot be read by this copy of the app. Save the sheet "
                         "as CSV and import that instead.")
    except Exception as exc:
        raise SheetError(f"That Excel file could not be opened.\n\n({exc})")


def _decode(raw: bytes) -> str:
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    head = raw[:2000]
    if head.count(b"\x00") > len(head) // 4:     # UTF-16 saved without its marker
        even, odd = head[0::2].count(b"\x00"), head[1::2].count(b"\x00")
        return raw.decode("utf-16-le" if odd >= even else "utf-16-be", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_table(path: str, sheet: str | None = None, max_rows: int | None = None,
               with_numbers: bool = False):
    """Read a spreadsheet. Returns (headers, rows): the first non-empty row is
    the headings; rows are lists of text, padded to the headings' length.
    Completely empty rows and columns without a heading or any content are dropped.
    with_numbers=True returns (headers, rows, numbers) where numbers[i] is the
    row number of rows[i] as the person sees it in their spreadsheet."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        raise SheetError("Old-style .xls files cannot be read. Open it in Excel and use "
                         "Save As to make an .xlsx or .csv file, then import that.")
    grid: list[list[str]] = []
    if ext in EXCEL:
        try:
            import openpyxl
        except ImportError:
            raise SheetError("Excel files cannot be read by this copy of the app. Save the "
                             "sheet as CSV and import that instead.")
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise SheetError(f"That Excel file could not be opened.\n\n({exc})")
        try:
            names = _visible_first(wb)
            if not names:
                raise SheetError("That Excel file has no worksheets in it.")
            ws = wb[sheet if sheet in names else names[0]]
            # some programs write a wrong size into the file; work it out instead
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            for row in ws.iter_rows():
                grid.append([_excel_cell(c) for c in row])
                if max_rows is not None and len(grid) > max_rows + 50:
                    break
        except SheetError:
            raise
        except Exception as exc:
            raise SheetError(f"That Excel file could not be read.\n\n({exc})")
        finally:
            wb.close()
    else:
        try:
            with open(path, "rb") as fh:
                raw = fh.read()
        except OSError as exc:
            raise SheetError(f"That file could not be read.\n\n({exc})")
        if raw[:4] in (b"PK\x03\x04", b"\xd0\xcf\x11\xe0", b"%PDF"):
            raise SheetError("That is not a CSV file, whatever its name says. If it is a "
                             "spreadsheet, open it and use Save As to make an .xlsx or .csv "
                             "file, then import that.")
        text = _decode(raw).replace("\x00", "")
        del raw
        if sum(1 for ch in text[:4000] if ch < " " and ch not in "\t\r\n\f") > len(text[:4000]) // 50:
            raise SheetError("That file does not hold a table of text, so it cannot be "
                             "imported. Save the information as a .csv or .xlsx file first.")
        m = re.match(r"sep=(.)\r?\n", text)      # Excel's own note of the separator
        if m:
            text = text[m.end():]
        sample = text[:20000]
        delim = ","
        if m:
            delim = m.group(1)
        else:
            try:
                delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                first = sample.split("\n", 1)[0]
                delim = max(",;\t|", key=first.count) if first else ","
        try:
            for row in csv.reader(io.StringIO(text, newline=""), delimiter=delim):
                grid.append([c.strip() for c in row])
                if max_rows is not None and len(grid) > max_rows + 50:
                    break
        except csv.Error as exc:
            if "field larger" in str(exc):
                raise SheetError("One cell in that file is far too long to be real. Usually "
                                 "this is a quotation mark (\") that was opened and never "
                                 f"closed, near row {len(grid) + 1}. Open the file, put that "
                                 "right, and try again.")
            raise SheetError(f"That file is not laid out as a table.\n\n({exc})")
    skipped = 0
    while grid and not any(grid[0]):
        grid.pop(0)
        skipped += 1
    if not grid:
        raise SheetError("That file is empty.")
    headers = grid[0]
    numbered = [(skipped + i + 2, r) for i, r in enumerate(grid[1:]) if any(r)]
    numbers = [n for n, _r in numbered]
    body = [r for _n, r in numbered]
    width = max([len(headers)] + [len(r) for r in body])
    headers = headers + [""] * (width - len(headers))
    body = [r + [""] * (width - len(r)) for r in body]
    keep = [i for i in range(width) if headers[i] or any(r[i] for r in body)]
    seen: dict[str, int] = {}
    out_headers = []
    for i in keep:
        h = headers[i] or f"Column {i + 1}"
        n = seen.get(h.lower(), 0) + 1
        seen[h.lower()] = n
        out_headers.append(h if n == 1 else f"{h} ({n})")
    rows = [[r[i] for i in keep] for r in body]
    if max_rows is not None:
        rows, numbers = rows[:max_rows], numbers[:max_rows]
    if with_numbers:
        return out_headers, rows, numbers
    return out_headers, rows
